"""DOCX Loader using python-docx."""
import re
from pathlib import Path
from typing import List, Dict, Any
from docx import Document


def clean_text(text: str) -> str:
    """Normalize whitespace."""
    if not text:
        return ""
    text = text.replace("\u00a0", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


def load_docx(file_path: Path, filename: str) -> List[Dict[str, Any]]:
    """
    Extracts structured paragraphs and tables from a DOCX file.
    """
    try:
        doc = Document(str(file_path))
    except Exception as e:
        raise ValueError(f"Failed to read Word document '{filename}': {str(e)}")

    blocks: List[Dict[str, Any]] = []
    current_heading = "General"
    current_paragraphs: List[str] = []

    def flush_section():
        nonlocal current_paragraphs
        if current_paragraphs:
            combined = "\n\n".join(current_paragraphs)
            cleaned = clean_text(combined)
            if cleaned:
                blocks.append({
                    "text": cleaned,
                    "page": None,
                    "section": current_heading,
                    "metadata": {
                        "section": current_heading
                    }
                })
            current_paragraphs = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        # Check if style is a heading
        style_name = para.style.name.lower() if para.style and para.style.name else ""
        if "heading" in style_name:
            flush_section()
            current_heading = text
        else:
            current_paragraphs.append(text)

    flush_section()

    # Also extract tables if present
    for table_idx, table in enumerate(doc.tables):
        rows_text = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows_text.append(" | ".join(cells))
        if rows_text:
            cleaned_table = "\n".join(rows_text)
            blocks.append({
                "text": cleaned_table,
                "page": None,
                "section": f"Table {table_idx + 1}",
                "metadata": {
                    "section": f"Table {table_idx + 1}"
                }
            })

    if not blocks:
        raise ValueError(f"No extractable text or content found in Word document '{filename}'.")

    return blocks
